import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

a = pd.read_csv('d/649.csv'); b = pd.read_csv('d/LOTTOMAX.csv')
a = a[a['DRAW DATE'] >= '2004-01-01']
bm = b[b['SEQUENCE NUMBER'] == 0]          # main Lotto Max draw
bx = b[b['SEQUENCE NUMBER'] > 0]           # MaxMillions draws
a.to_csv('/mnt/user-data/outputs/lotto649_2004_2026.csv', index=False)
bm.to_csv('/mnt/user-data/outputs/lottomax_main_2009_2026.csv', index=False)
bx.to_csv('/mnt/user-data/outputs/lottomax_maxmillions.csv', index=False)

F = Font(name='Arial'); H = Font(name='Arial', bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F4E78')
wb = Workbook(); wb.remove(wb.active)

def sheet(name, df, cols, heads):
    ws = wb.create_sheet(name); ws.append(heads)
    for r in df[cols].itertuples(index=False): ws.append(list(r))
    for c in ws[1]: c.font, c.fill, c.alignment = H, FILL, Alignment(horizontal='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i in range(1, len(heads)+1): ws.column_dimensions[get_column_letter(i)].width = 12
    return ws

n6 = [f'NUMBER DRAWN {i}' for i in range(1,7)]; n7 = [f'NUMBER DRAWN {i}' for i in range(1,8)]
s1 = sheet('Lotto 6-49', a, ['DRAW NUMBER','DRAW DATE']+n6+['BONUS NUMBER'],
           ['Tirage #','Date','N1','N2','N3','N4','N5','N6','Complémentaire'])
s2 = sheet('Lotto Max', bm, ['DRAW NUMBER','DRAW DATE']+n7+['BONUS NUMBER'],
           ['Tirage #','Date','N1','N2','N3','N4','N5','N6','N7','Complémentaire'])
s3 = sheet('MaxMillions', bx, ['DRAW NUMBER','SEQUENCE NUMBER','DRAW DATE']+n7,
           ['Tirage #','Séquence','Date','N1','N2','N3','N4','N5','N6','N7'])

# Frequency sheet (live formulas)
fq = wb.create_sheet('Fréquences')
fq.append(['Numéro','6/49 (principaux)','6/49 (complém.)','Lotto Max (principaux)','Lotto Max (complém.)'])
for c in fq[1]: c.font, c.fill = H, FILL
e1, e2 = s1.max_row, s2.max_row
for n in range(1, 51):
    r = n + 1
    fq.append([n,
      f"=IF(A{r}>49,\"\",COUNTIF('Lotto 6-49'!$C$2:$H${e1},A{r}))",
      f"=IF(A{r}>49,\"\",COUNTIF('Lotto 6-49'!$I$2:$I${e1},A{r}))",
      f"=COUNTIF('Lotto Max'!$C$2:$I${e2},A{r})",
      f"=COUNTIF('Lotto Max'!$J$2:$J${e2},A{r})"])
for row in fq.iter_rows(min_row=2):
    for c in row: c.font = F
for i in range(1,6): fq.column_dimensions[get_column_letter(i)].width = 22

src = wb.create_sheet('Source')
for line in ['Source officielle : BCLC PlayNow — fichiers « downloadable numbers » (649.zip, LOTTOMAX.zip)',
             'https://www.playnow.com/resources/documents/downloadable-numbers/649.zip',
             'https://www.playnow.com/resources/documents/downloadable-numbers/LOTTOMAX.zip',
             'Lotto 6/49 et Lotto Max sont des jeux nationaux : mêmes numéros au Québec.',
             'Lotto 6/49 : depuis 2004-01-03. Lotto Max : depuis le lancement 2009-09-25.',
             'Non inclus : Québec 49, Québec Max, Grande Vie, Extra, Banco (jeux propres à Loto-Québec).',
             'Données téléchargées le 2026-09-26. Résultats non officiels — valider auprès de Loto-Québec.']:
    src.append([line]); src.cell(src.max_row,1).font = F
src.column_dimensions['A'].width = 100
wb.save('/mnt/user-data/outputs/loto_quebec_resultats_2004_2026.xlsx')
print(len(a), len(bm), len(bx))
