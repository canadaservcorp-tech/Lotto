"""Lotto Max — répétition et fréquence (ère 7/50 : 2019-05-14 au 2026-04-10). Cotes : jeu actuel 7/52."""
import itertools, math
from collections import Counter
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

CUTOFF, N, K = '2019-05-14', 50, 7
b = pd.read_csv('d/LOTTOMAX.csv')
m = b[(b['SEQUENCE NUMBER'] == 0) & (b['DRAW DATE'] >= CUTOFF)
      & (b['DRAW DATE'] < '2026-04-14')].sort_values('DRAW DATE')   # 7/52 depuis le 2026-04-14
cols = [f'NUMBER DRAWN {i}' for i in range(1, 8)]
draws = [sorted(r) for r in m[cols].values.tolist()]
D = len(draws); last = D + 1                          # last data row in sheet

F, B = Font(name='Arial'), Font(name='Arial', bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F4E78'); NOTE = Font(name='Arial', italic=True, color='666666')
wb = Workbook(); wb.remove(wb.active)

def header(ws, heads, widths=14):
    ws.append(heads)
    for c in ws[1]: c.font, c.fill, c.alignment = B, FILL, Alignment(horizontal='center', wrap_text=True)
    for i in range(1, len(heads) + 1): ws.column_dimensions[L(i)].width = widths
    ws.freeze_panes = 'A2'

def fontall(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font != NOTE: c.font = F

# 1. Tirages + répétitions vs tirage précédent (formule)
t = wb.create_sheet('Tirages')
header(t, ['Tirage #', 'Date', 'N1', 'N2', 'N3', 'N4', 'N5', 'N6', 'N7', 'Compl.', 'Répétés du tirage précédent'])
for i, (_, r) in enumerate(m.iterrows()):
    row = i + 2
    rep = '' if row == 2 else f'=SUMPRODUCT(COUNTIF(C{row-1}:I{row-1},C{row}:I{row}))'
    t.append([r['DRAW NUMBER'], r['DRAW DATE']] + draws[i] + [r['BONUS NUMBER'], rep])
t.auto_filter.ref = t.dimensions; fontall(t)
RNG = f"Tirages!$C$2:$I${last}"

# 2. Fréquences
f = wb.create_sheet('Fréquences')
header(f, ['Numéro', 'Sorties', 'Attendu', 'Écart', 'Ratio obs/att', '% tirages observé',
           'Probabilité théorique', 'Dernière ligne', 'Tirages depuis (retard)', 'Écart-type (z)'], 15)
for n in range(1, N + 1):
    r = n + 1
    f.append([n, f'=COUNTIF({RNG},A{r})', f'=$M$1*$M$2', f'=B{r}-C{r}', f'=B{r}/C{r}',
              f'=B{r}/$M$1', '=$M$2',
              f'=SUMPRODUCT(MAX(({RNG}=A{r})*ROW({RNG})))', f'={last}-H{r}',
              f'=(B{r}-C{r})/SQRT($M$1*$M$2*(1-$M$2))'])
for r in range(2, N + 2):
    f[f'E{r}'].number_format = '0.00'; f[f'F{r}'].number_format = '0.0%'
    f[f'G{r}'].number_format = '0.0%'; f[f'J{r}'].number_format = '0.00'
f['L1'], f['M1'] = 'Nb tirages', f'=COUNT(Tirages!$A$2:$A${last})'
f['L2'], f['M2'] = 'Prob. par tirage', f'={K}/{N}'
f['A54'] = 'z entre -2 et +2 = variation normale du hasard. Chaque tirage est indépendant : le passé ne prédit pas le prochain.'
f['A54'].font = NOTE; fontall(f)
ch = BarChart(); ch.title = 'Sorties observées vs attendues'; ch.height, ch.width = 9, 28
ch.add_data(Reference(f, min_col=2, max_col=3, min_row=1, max_row=N + 1), titles_from_data=True)
ch.set_categories(Reference(f, min_col=1, min_row=2, max_row=N + 1)); f.add_chart(ch, 'L5')

# 3. Répétitions : distribution observée vs loi hypergéométrique
p = wb.create_sheet('Répétitions')
header(p, ['Numéros répétés du tirage précédent', 'Fois observé', '% observé',
           'Probabilité théorique', 'Fois attendu'], 20)
for k in range(K + 1):
    r = k + 2
    p.append([k, f'=COUNTIF(Tirages!$K$3:$K${last},A{r})', f'=B{r}/SUM($B$2:$B$9)',
              f'=COMBIN({K},A{r})*COMBIN({N-K},{K}-A{r})/COMBIN({N},{K})', f'=D{r}*SUM($B$2:$B$9)'])
    p[f'C{r}'].number_format = p[f'D{r}'].number_format = '0.00%'; p[f'E{r}'].number_format = '0.0'
p['A11'], p['B11'] = 'Moyenne observée', f'=AVERAGE(Tirages!$K$3:$K${last})'
p['A12'], p['B12'] = 'Moyenne théorique', f'={K}*{K}/{N}'
p['A14'] = 'Théorie : hypergéométrique C(7,k)·C(43,7−k)/C(50,7).'; p['A14'].font = NOTE
fontall(p)

# 4. Paires les plus fréquentes (calcul Python, valeurs statiques)
pc = Counter(pr for d in draws for pr in itertools.combinations(d, 2))
q = wb.create_sheet('Paires')
header(q, ['Numéro A', 'Numéro B', 'Fois ensemble', 'Attendu', 'Ratio'])
exp_pair = D * (K * (K - 1)) / (N * (N - 1))
for (x, y), c in pc.most_common(40):
    q.append([x, y, c, round(exp_pair, 2), round(c / exp_pair, 2)])
q['G1'] = f'Calcul Python sur {D} tirages (valeurs figées, relancer le script pour mettre à jour).'
q['G1'].font = NOTE; fontall(q)

# 5. Combinaisons répétées / chevauchement max
combos = Counter(tuple(d) for d in draws)
dup = [c for c, v in combos.items() if v > 1]
best = max(((len(set(a) & set(b)), i, j) for (i, a), (j, b)
            in itertools.combinations(enumerate(draws), 2)))
c5 = wb.create_sheet('Combinaisons')
for line in [f'Tirages analysés : {D} (7/50, {CUTOFF} au 2026-04-10)',
             f'Combinaisons complètes répétées : {len(dup)}',
             f'Plus grand chevauchement entre 2 tirages : {best[0]} numéros '
             f'({m.iloc[best[1]]["DRAW DATE"]} et {m.iloc[best[2]]["DRAW DATE"]})',
             f'Combinaisons possibles : {math.comb(N, K):,}']:
    c5.append([line])
c5.column_dimensions['A'].width = 90; fontall(c5)

# 6. Probabilités par catégorie de lot (formules COMBIN)
g = wb.create_sheet('Probabilités lots')
header(g, ['Catégorie', 'Probabilité (1 sur …)', 'Probabilité %'], 22)
NC = 52   # jeu actuel depuis le 2026-04-14
C = f'COMBIN({NC},7)'; R = NC - 8
tiers = [('7/7', f'={C}'),
         ('6/7 + compl.', f'={C}/COMBIN(7,6)'),
         ('6/7', f'={C}/(COMBIN(7,6)*COMBIN({R},1))'),
         ('5/7 + compl.', f'={C}/(COMBIN(7,5)*COMBIN({R},1))'),
         ('5/7', f'={C}/(COMBIN(7,5)*COMBIN({R},2))'),
         ('4/7 + compl.', f'={C}/(COMBIN(7,4)*COMBIN({R},2))'),
         ('4/7', f'={C}/(COMBIN(7,4)*COMBIN({R},3))'),
         ('3/7 + compl.', f'={C}/(COMBIN(7,3)*COMBIN({R},3))'),
         ('3/7', f'={C}/(COMBIN(7,3)*COMBIN({R},4))')]
for i, (name, fx) in enumerate(tiers):
    r = i + 2; g.append([name, fx, f'=1/B{r}'])
    g[f'B{r}'].number_format = '#,##0'; g[f'C{r}'].number_format = '0.000000%'
g['A12'] = 'Jeu actuel 7/52, par ligne. Un billet de 6 $ = 4 lignes.'; g['A12'].font = NOTE
fontall(g)

wb.save('lotto_max_analyse.xlsx')
print(D, len(dup), best[0])
